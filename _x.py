import sys
from pptx import Presentation
from pptx.util import Emu
p=Presentation(sys.argv[1])
out=[]
for i,s in enumerate(p.slides,1):
    out.append(f"===== SLIDE {i} =====")
    def walk(shapes,pfx=""):
        for sh in shapes:
            if sh.shape_type==6:
                walk(sh.shapes,pfx+"  ")
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(pfx+sh.text_frame.text.replace("\n"," | "))
            if sh.has_table:
                for r in sh.table.rows:
                    out.append(pfx+"TBL: "+" | ".join(c.text.replace("\n"," ") for c in r.cells))
    walk(s.shapes)
print("\n".join(out))
